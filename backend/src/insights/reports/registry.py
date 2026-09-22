"""Report definitions and the registry the API serves them from.

Adding a report means registering one more :class:`ReportDefinition`; the HTTP layer stays
untouched (design spec §3.6 / §8).
"""

from __future__ import annotations

import csv
import io
import sqlite3
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date
from typing import Any, Literal

ColumnType = Literal["string", "number", "date", "month", "credits", "currency"]
ParamType = Literal["date", "string", "number"]


class ReportError(Exception):
    """Invalid report parameters (surfaced as HTTP 400)."""


@dataclass(frozen=True)
class ReportParam:
    name: str
    type: ParamType
    required: bool = False
    description: str = ""

    def as_json(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "type": self.type,
            "required": self.required,
            "description": self.description,
        }


@dataclass(frozen=True)
class ReportColumn:
    key: str
    label: str
    type: ColumnType

    def as_json(self) -> dict[str, str]:
        return {"key": self.key, "label": self.label, "type": self.type}


@dataclass
class ReportContext:
    """What a report receives: an open store connection, the range and extra params."""

    conn: sqlite3.Connection
    date_from: date
    date_to: date
    params: dict[str, str] = field(default_factory=dict)

    def require(self, name: str) -> str:
        value = self.params.get(name)
        if not value:
            raise ReportError(f"missing required parameter '{name}'")
        return value


@dataclass
class ReportResult:
    """Columns, rows, an optional Total row (keyed by column key) and free-form metadata."""

    columns: list[ReportColumn]
    rows: list[dict[str, Any]]
    totals: dict[str, Any] | None = None
    meta: dict[str, Any] | None = None


ReportRunner = Callable[[ReportContext], ReportResult]


@dataclass(frozen=True)
class ReportDefinition:
    key: str
    title: str
    description: str
    params: tuple[ReportParam, ...]
    run: ReportRunner

    def describe(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "title": self.title,
            "description": self.description,
            "params": [p.as_json() for p in self.params],
        }


class ReportRegistry:
    """Ordered registry of report definitions keyed by ``key``."""

    def __init__(self) -> None:
        self._defs: dict[str, ReportDefinition] = {}

    def register(self, definition: ReportDefinition) -> ReportDefinition:
        self._defs[definition.key] = definition
        return definition

    def get(self, key: str) -> ReportDefinition | None:
        return self._defs.get(key)

    def list(self) -> list[ReportDefinition]:
        return list(self._defs.values())


registry = ReportRegistry()
"""Process-wide registry; the only intentional piece of global mutable state."""


def result_to_json(
    definition: ReportDefinition, ctx: ReportContext, result: ReportResult
) -> dict[str, Any]:
    return {
        "key": definition.key,
        "title": definition.title,
        "from": ctx.date_from.isoformat(),
        "to": ctx.date_to.isoformat(),
        "columns": [c.as_json() for c in result.columns],
        "rows": result.rows,
        "totals": result.totals,
        "meta": result.meta,
    }


def result_to_csv(result: ReportResult) -> str:
    """Serialise rows (and a trailing totals row when present) as CSV text."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([c.label for c in result.columns])
    for row in result.rows:
        writer.writerow([_cell(row.get(c.key)) for c in result.columns])
    if result.totals:
        writer.writerow([_cell(result.totals.get(c.key)) for c in result.columns])
    return buffer.getvalue()


def result_to_xlsx(result: ReportResult, title: str = "Report") -> bytes:
    """Serialise the report as a single-sheet workbook: bold header, numeric cells, Total row."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = title[:31] or "Report"
    sheet.append([c.label for c in result.columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    numeric = {c.key for c in result.columns if c.type in ("number", "credits", "currency")}
    for row in result.rows:
        sheet.append([_xlsx_cell(row.get(c.key), c.key in numeric) for c in result.columns])
    if result.totals:
        sheet.append(
            [_xlsx_cell(result.totals.get(c.key), c.key in numeric) for c in result.columns]
        )
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
    for index, column in enumerate(result.columns, start=1):
        width = max(len(column.label), 10)
        for row in sheet.iter_rows(min_col=index, max_col=index, min_row=2):
            for cell in row:
                if cell.value is not None:
                    width = max(width, min(len(str(cell.value)), 48))
        sheet.column_dimensions[get_column_letter(index)].width = width + 2
        if column.key in numeric:
            for row in sheet.iter_rows(min_col=index, max_col=index, min_row=2):
                for cell in row:
                    cell.number_format = "#,##0.00"
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _xlsx_cell(value: Any, numeric: bool) -> Any:
    if value is None:
        return None
    if numeric:
        return round(float(value), 4) if isinstance(value, (int, float)) else value
    return value


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".")
    return str(value)
